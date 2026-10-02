"""Word import: read any .docx, let the model map it onto our layouts.

Two steps, kept apart on purpose:

1. `outline()` reads the document without any model: every paragraph and table
   in order, with the clues a human would use to see structure (style name,
   heading level, size, bold, list item, shaded box). This part is plain code,
   so it is cheap, testable and identical for every provider.
2. `map_document()` hands that outline to the model, which decides what is the
   masthead, which lines are chapter titles, what belongs in a highlight box,
   and says where it was unsure. The text itself is copied, never rewritten.
"""
from __future__ import annotations

import io
import json
import re
from typing import Literal

from docx import Document
from docx.oxml.ns import qn
from pydantic import BaseModel, Field, ValidationError

from .providers import GenerationError, Provider
from .schemas import ICONS

MAX_OUTLINE_CHARS = 60_000
MAX_CHAPTERS = 10  # as in the editor


# ---------- 1. reading the document ----------

def _para_info(p) -> dict:
    text = p.text.strip()
    style = p.style.name if p.style is not None else ""
    pPr = p._p.pPr
    is_list = bool(pPr is not None and pPr.find(qn("w:numPr")) is not None) or "List" in style
    outline_lvl = None
    if pPr is not None and pPr.find(qn("w:outlineLvl")) is not None:
        outline_lvl = int(pPr.find(qn("w:outlineLvl")).get(qn("w:val")))
    m = re.match(r"Heading (\d)", style)
    if m:
        outline_lvl = int(m.group(1)) - 1
    runs = [r for r in p.runs if r.text.strip()]
    chars = sum(len(r.text) for r in runs) or 1
    bold = sum(len(r.text) for r in runs if r.bold) / chars
    sizes = [r.font.size.pt for r in runs if r.font.size is not None]
    if not sizes and p.style is not None and p.style.font.size is not None:
        sizes = [p.style.font.size.pt]
    caps = bool(text) and text.upper() == text and any(c.isalpha() for c in text)
    return {
        "text": text,
        "style": style,
        "level": outline_lvl,
        "list": is_list,
        "bold": round(bold, 2),
        "size": max(sizes) if sizes else None,
        "caps": caps,
    }


def _cell_fill(cell) -> str | None:
    shd = cell._tc.find(qn("w:tcPr") + "/" + qn("w:shd"))
    fill = shd.get(qn("w:fill")) if shd is not None else None
    return fill if fill and fill.lower() not in ("auto", "ffffff") else None


def outline(data: bytes) -> tuple[str, int]:
    """A numbered, line-per-element description of the document for the model.

    Returns (outline, number_of_elements)."""
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 - any parse failure means "not a docx"
        raise GenerationError("That file could not be read as a Word document (.docx).") from e

    lines: list[str] = []

    def add(kind: str, info: dict | None, text: str) -> None:
        tags = [kind]
        if info:
            if info["style"] and info["style"] != "Normal":
                tags.append(f"style={info['style']}")
            if info["level"] is not None:
                tags.append(f"heading{info['level'] + 1}")
            if info["size"]:
                tags.append(f"{info['size']:g}pt")
            if info["bold"] >= 0.8:
                tags.append("bold")
            if info["caps"]:
                tags.append("caps")
            if info["list"]:
                tags.append("bullet")
        lines.append(f"[{len(lines) + 1}] {' '.join(tags)}: {text}")

    # Headers often carry the publication name.
    seen: set[str] = set()
    for section in doc.sections:
        for p in section.header.paragraphs:
            t = p.text.strip()
            if t and t not in seen:
                seen.add(t)
                add("page-header", None, t)

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            p = Paragraph(child, doc)
            info = _para_info(p)
            if info["text"]:
                add("p", info, info["text"])
            elif child.findall(".//" + qn("w:drawing")):
                add("image", None, "(picture, not imported)")
        elif child.tag == qn("w:tbl"):
            from docx.table import Table

            table = Table(child, doc)
            cells = [c for row in table.rows for c in row.cells]
            unique = list(dict.fromkeys(cells))  # merged cells repeat
            if len(table.rows) == 1 and len(table.columns) <= 2:
                # Usually a box or a two-column layout trick, not data.
                for c in unique:
                    fill = _cell_fill(c)
                    kind = f"box(fill=#{fill})" if fill else "box"
                    for p in c.paragraphs:
                        info = _para_info(p)
                        if info["text"]:
                            add(kind, info, info["text"])
            else:
                for row in table.rows:
                    row_cells = list(dict.fromkeys(row.cells))
                    text = " | ".join(c.text.strip().replace("\n", " / ") for c in row_cells)
                    if text.strip(" |"):
                        add("table-row", None, text)

    text = "\n".join(lines)
    if len(text) > MAX_OUTLINE_CHARS:
        text = text[:MAX_OUTLINE_CHARS] + "\n[... document truncated ...]"
    return text, len(lines)


# ---------- 2. mapping it onto our layouts ----------

class ImportBox(BaseModel):
    type: Literal["action_box", "info_box"] = "info_box"
    title: str = ""
    text: str = ""


class ImportChapter(BaseModel):
    heading: str = ""
    icon: str = ""
    text: str = ""
    boxes: list[ImportBox] = Field(default_factory=list)  # tolerated; folded into text


class ImportCard(BaseModel):
    icon: str = "info"
    title: str = ""
    subtitle: str = ""
    text: str = ""


class ImportNews(BaseModel):
    icon: str = "smile"
    label: str = ""
    text: str = ""


class ImportMasthead(BaseModel):
    header_kicker: str = ""
    doc_type: str = ""
    doc_issue: str = ""
    doc_headline: str = ""
    publication_date: str = ""   # YYYY-MM-DD or empty


class ImportFlag(BaseModel):
    where: str
    note: str


class ImportResult(BaseModel):
    layout: Literal["standard", "one_pager"] = "standard"
    layout_reason: str = ""
    masthead: ImportMasthead = Field(default_factory=ImportMasthead)
    chapters: list[ImportChapter] = Field(default_factory=list)
    cards: list[ImportCard] = Field(default_factory=list)
    news: list[ImportNews] = Field(default_factory=list)
    flags: list[ImportFlag] = Field(default_factory=list)


IMPORT_SYSTEM = """You restructure existing internal staff publications of Helvetic
Airways, a Swiss regional airline, into the airline's newsletter layouts. You
move text into the right places; you never write, shorten or improve it.
Answer with one JSON object and nothing else."""


def import_prompt(outline_text: str, layout: str, icons: list[dict], kickers: list[str]) -> str:
    icon_list = ", ".join(f"{i['v']} ({i['label']})" for i in icons if i.get("v") in ICONS)
    if layout == "auto":
        layout_rule = (
            'Choose "layout": "one_pager" if the content is a handful of short items that '
            'fit one A4 page as up to 4 cards and up to 3 one-to-three-line news items; '
            'otherwise "standard". Explain the choice in one sentence in "layout_reason".'
        )
    else:
        layout_rule = (
            f'The editor chose "layout": "{layout}". Use it. If the content does not suit '
            'it (for example too much for one page), still use it and add a flag with '
            'where "layout" saying so.'
        )
    return f"""Below is an outline of a Word document. Each line is one paragraph or
table cell in reading order: [number] then clues (style, heading level, font size,
bold, caps, bullet, box with fill colour, table row), then the text.

Map it onto this JSON structure:

{{
  "layout": "standard" | "one_pager",
  "layout_reason": "...",
  "masthead": {{
    "header_kicker": "publication name, e.g. one of: {'; '.join(kickers)}",
    "doc_type": "Newsletter" | "Bulletin",
    "doc_issue": "issue, e.g. '10 2026' or 'Q3 2026'",
    "doc_headline": "the edition's headline or subtitle, if any",
    "publication_date": "YYYY-MM-DD, or empty if none is stated"
  }},
  "chapters": [   // only for "standard"
    {{"heading": "...", "icon": "<icon key or empty>", "text": "..."}}
  ],
  "cards": [      // only for "one_pager", at most 4
    {{"icon": "<icon key>", "title": "...", "subtitle": "...", "text": "..."}}
  ],
  "news": [       // only for "one_pager", at most 3, each one to three short lines
    {{"icon": "<icon key>", "label": "...", "text": "..."}}
  ],
  "flags": [{{"where": "...", "note": "..."}}]
}}

Rules:
- {layout_rule}
- Copy text exactly as it is in the document: same words, spelling and order.
  Do not summarise, shorten, correct or add anything. Keep paragraph breaks as
  line breaks. Write bullet items as lines starting with "- ".
- Chapter titles and card titles are the document's own headings.
- A chapter can hold highlight boxes, written inside its "text" as one line each,
  at the place they appear: "[ACTION] Title: text" (red, for something the reader
  must do or a deadline) or "[INFO] Title: text" (blue, for background). Use them
  only for content visibly set apart in the source, such as a shaded box or a
  "Note:" / "Important:" paragraph. If the source gives the box no title, use a
  short one from its own first words (e.g. "Action" from "Action: ...").
- Leave out repeated page headers, footers, page numbers, tables of contents
  and the masthead lines themselves once used for "masthead".
- Pictures cannot be imported; if a picture seems to carry information, flag it.
- Pick a fitting icon for every chapter, card and news item from: {icon_list}.
- Flag everything you had to guess or could not place, so the editor can check
  it. "where" names the field: "layout", "masthead.<field>",
  "chapters.<index>.<heading|icon|text>", "cards.<index>.<title|subtitle|icon|text>",
  "news.<index>.<label|icon|text>" (indexes from 0). "note" is one short sentence
  for the editor, e.g. "Guessed that this bold line is a chapter title." or
  "No issue number found in the document." Do not flag icons you chose with
  confidence, and do not flag things that were clear.

<document>
{outline_text}
</document>"""


def _parse(raw: str) -> ImportResult:
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < 0:
        raise GenerationError("The AI did not return a usable structure. Please try again.")
    try:
        return ImportResult.model_validate(json.loads(text[start:end + 1]))
    except (json.JSONDecodeError, ValidationError) as e:
        raise GenerationError("The AI returned an incomplete structure. Please try again.") from e


async def map_document(
    data: bytes, layout: str, provider: Provider, icons: list[dict], kickers: list[str]
) -> tuple[ImportResult, int]:
    if provider.name == "manual":
        raise GenerationError("Importing a Word document needs an AI writer. Pick one at the top.")
    outline_text, count = outline(data)
    if not count:
        raise GenerationError("That document has no text to import.")
    raw = await provider.complete(
        IMPORT_SYSTEM, import_prompt(outline_text, layout, icons, kickers), max_tokens=16000
    )
    result = _parse(raw)
    if layout != "auto":
        result.layout = layout  # the editor's choice stands
    # Keep within what the layouts can hold; say so rather than drop silently.
    if result.layout == "one_pager":
        if len(result.cards) > 4:
            result.flags.append(ImportFlag(
                where="layout",
                note=f"The document had {len(result.cards)} card-sized items; only the first 4 fit.",
            ))
            result.cards = result.cards[:4]
        if len(result.news) > 3:
            result.flags.append(ImportFlag(
                where="layout",
                note=f"The document had {len(result.news)} short news; only the first 3 fit.",
            ))
            result.news = result.news[:3]
    if len(result.chapters) > MAX_CHAPTERS:
        result.flags.append(ImportFlag(
            where="layout",
            note=f"The document had {len(result.chapters)} chapters; only the first {MAX_CHAPTERS} fit.",
        ))
        result.chapters = result.chapters[:MAX_CHAPTERS]
    for c in result.chapters:  # boxes belong inline, as the editor shows them
        for b in c.boxes:
            tag = "ACTION" if b.type == "action_box" else "INFO"
            c.text = f"{c.text.rstrip()}\n[{tag}] {b.title.replace(':', ' -')}: {' '.join(b.text.split())}".strip()
        c.boxes = []
    for item in [*result.chapters, *result.cards, *result.news]:
        if item.icon and item.icon not in ICONS:
            item.icon = "info"
    return result, count

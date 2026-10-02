"""Renders the one-page bulletin: up to four cards plus up to three news rows.

Design lives in template_onepager.docx; this module only decides which named
style each piece gets and how the boxes are arranged. Two arrangement rules
matter:

* Cards sit two per row. An odd last card spans the full width instead of
  leaving a hole, so two cards fill a page as comfortably as four.
* Empty cards and empty news rows are dropped before layout, so an unused box
  costs nothing and the remaining boxes get the space.
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from ..budget import card_rows, fit_size, is_full_width
from ..schemas import Masthead, RenderedCard, RenderedNews
from . import docxutil as X

ASSETS = Path(__file__).resolve().parent.parent / "assets"
TEMPLATE = ASSETS / "template_onepager.docx"

RED = "E4032E"
CARD_BG = "F2F3F4"
CARD_LINE = "DDE0E2"

CONTENT_W = 9978          # A4 less the layout's 1.7 cm margins
GAP = 280
HALF = (CONTENT_W - GAP) // 2
NEWS_LABEL_W = 2850

# Row minimums: with one row of cards there is far more height to give it.
MIN_ROW_ONE = 7200
MIN_ROW_TWO = 3200


CARD_ICON_PT = 28
NEWS_ICON_PT = 24
ICON_GAP = 160            # twips between an icon and the text beside it
CARD_PAD = 160            # card cell's left/right margin, see _fill_card


def _icon_header(cell, icon: str, size_pt: int, width: int):
    """Put a borderless two-column table at the top of `cell`: the round icon
    on the left, and return the right-hand cell for the heading text.

    A side-by-side table rather than an inline picture, so a two-line heading
    sits beside the icon instead of wrapping underneath it.
    """
    icon_w = size_pt * 20 + ICON_GAP
    table = cell.add_table(rows=1, cols=2)
    # add_table appends after the cell's starting paragraph and adds an empty
    # one after itself; drop both so the header is the first thing in the cell.
    for p in cell.paragraphs:
        p._p.getparent().remove(p._p)
    X.table_no_borders(table)
    table.autofit = False
    X.fixed_columns(table, [icon_w, width - icon_w])
    icon_cell, text_cell = table.rows[0].cells
    for c in (icon_cell, text_cell):
        X.cell_margins(c, top=0, bottom=0, left=0, right=0)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    path = ASSETS / f"ic_{icon}.png"
    if path.exists():
        pic = icon_cell.paragraphs[0]
        pic.paragraph_format.space_after = Pt(0)
        pic.paragraph_format.space_before = Pt(0)
        pic.add_run().add_picture(str(path), width=Pt(size_pt), height=Pt(size_pt))
    return text_cell


def _ensure_trailing_paragraph(cell) -> None:
    """Word requires a cell to end with a paragraph, not a table."""
    if cell._tc[-1].tag != qn("w:p"):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)


def _fill_card(cell, card: RenderedCard, width: int, size: float) -> None:
    X.shade_cell(cell, CARD_BG)
    X.cell_borders(cell, left=CARD_LINE, bottom=CARD_LINE, size=4)
    # cell_borders clears the sides it is not given, so restore the full frame
    X.all_borders(cell, CARD_LINE, 4)
    X.cell_margins(cell, top=120, bottom=120, left=CARD_PAD, right=CARD_PAD)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    head = _icon_header(cell, card.icon, CARD_ICON_PT, width - 2 * CARD_PAD)
    # Headings stay one point above the body, as in the template (9 pt over 8 pt).
    title = head.paragraphs[0]
    X.set_style(title, "CardTitle")
    title.add_run(card.title.upper())
    X.set_size(title, size + 1)

    if card.subtitle.strip():
        sub = head.add_paragraph(card.subtitle)
        X.set_style(sub, "CardSubtitle")
        X.set_size(sub, size + 1)
    # Close the heading's last paragraph flush with the icon so it centres.
    head.paragraphs[-1].paragraph_format.space_after = Pt(0)
    # Breathing room between the heading row and the card's body text.
    gap = cell.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(5)

    bullet_numpr = getattr(cell, "_bullet_numpr", None)
    for block in card.blocks:
        if block.kind == "bullets":
            for k, item in enumerate(block.items):
                p = cell.add_paragraph(item)
                X.set_style(p, "CardBody")
                X.apply_numpr(p, bullet_numpr)
                X.set_size(p, size)
                # bullets sit close together; the list as a whole keeps the
                # paragraph gap after its last item
                if k < len(block.items) - 1:
                    p.paragraph_format.space_after = Pt(1)
        elif block.kind == "body":
            for line in [t for t in block.text.split("\n") if t.strip()]:
                p = cell.add_paragraph(line.strip())
                X.set_style(p, "CardBody")
                X.set_size(p, size)
        else:
            # A highlight box inside a card would fight the card itself, so the
            # box title becomes a sub-heading and its text ordinary body copy.
            if block.title:
                X.set_style(cell.add_paragraph(block.title), "CardSub")
            p = cell.add_paragraph(block.text)
            X.set_style(p, "CardBody")
            X.set_size(p, size)
    _ensure_trailing_paragraph(cell)


def _card_tables(doc, cards: list[RenderedCard], bullet_numpr, size: float):
    rows = card_rows(len(cards))
    min_height = MIN_ROW_ONE if rows == 1 else MIN_ROW_TWO
    tables = []
    i = 0
    while i < len(cards):
        full = is_full_width(i, len(cards))
        table = doc.add_table(rows=1, cols=1 if full else 3)
        X.table_no_borders(table)
        table.autofit = False
        row = table.rows[0]
        X.row_min_height(row, min_height)
        X.row_cannot_split(row)

        if full:
            X.fixed_columns(table, [CONTENT_W])
            cell = row.cells[0]
            cell._bullet_numpr = bullet_numpr
            _fill_card(cell, cards[i], CONTENT_W, size)
            i += 1
        else:
            X.fixed_columns(table, [HALF, GAP, HALF])
            left, gap, right = row.cells
            X.clear_borders(gap)
            left._bullet_numpr = bullet_numpr
            right._bullet_numpr = bullet_numpr
            _fill_card(left, cards[i], HALF, size)
            _fill_card(right, cards[i + 1], HALF, size)
            i += 2
        tables.append(table)
        doc.add_paragraph()  # gap below the row
    return tables


def _news_table(doc, news: list[RenderedNews], size: float) -> None:
    if not news:
        return
    table = doc.add_table(rows=len(news), cols=2)
    X.table_no_borders(table)
    table.autofit = False
    X.fixed_columns(table, [NEWS_LABEL_W, CONTENT_W - NEWS_LABEL_W])
    for idx, item in enumerate(news):
        row = table.rows[idx]
        X.row_cannot_split(row)
        label_cell, text_cell = row.cells
        for cell, right_margin in ((label_cell, 180), (text_cell, 0)):
            X.cell_borders(cell, bottom=CARD_LINE, size=4, top=CARD_LINE if idx == 0 else None)
            X.cell_margins(cell, top=62, bottom=62, left=0, right=right_margin)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

        head = _icon_header(label_cell, item.icon, NEWS_ICON_PT, NEWS_LABEL_W - 180)
        label = head.paragraphs[0]
        X.set_style(label, "NewsLabel")
        label.add_run(item.label.upper())
        X.set_size(label, size + 1)
        _ensure_trailing_paragraph(label_cell)

        first = True
        for line in [l for l in item.lines if l.strip()]:
            p = text_cell.paragraphs[0] if first else text_cell.add_paragraph()
            if first:
                p.text = line.strip()
                first = False
            else:
                p.text = line.strip()
            X.set_style(p, "NewsBody")
            X.set_size(p, size)


def _paragraphs(card: RenderedCard) -> list[tuple[str, bool]]:
    """A card's body as (text, is_bullet) pairs, for fit_size."""
    out: list[tuple[str, bool]] = []
    for block in card.blocks:
        if block.kind == "bullets":
            out += [(item, True) for item in block.items]
        elif block.kind == "body":
            out += [(t.strip(), False) for t in block.text.split("\n") if t.strip()]
        else:
            out += [(block.title, False), (block.text, False)]
    return out


def render_one_pager(
    masthead: Masthead, cards: list[RenderedCard], news: list[RenderedNews]
) -> bytes:
    doc = Document(str(TEMPLATE))
    bullet_numpr = X.find_bullet_numpr(doc)
    X.clear_body(doc)

    mapping = {
        "KICKER": masthead.header_kicker,
        "ISSUED_BY": masthead.footer_issued_by,
        "REVISION": masthead.footer_revision,
        "PUB_DATE": masthead.publication_date.strftime("%d.%m.%Y"),
    }
    for section in doc.sections:
        X.replace_in_part(section.header, mapping)
        X.replace_in_part(section.footer, mapping)

    title = doc.add_paragraph()
    X.set_style(title, "DocTitle")
    title.add_run(f"{masthead.doc_type} \u2013 ")
    title.add_run(masthead.doc_issue).font.color.rgb = RGBColor.from_string(RED)

    X.set_style(doc.add_paragraph(masthead.doc_headline), "DocHeadline")
    X.set_style(
        doc.add_paragraph(masthead.publication_date.strftime("%d.%m.%Y")), "DocDate"
    )

    usable_cards = [c for c in cards if c.title.strip() or c.blocks]
    usable_news = [n for n in news if n.label.strip() or any(l.strip() for l in n.lines)]

    size = fit_size(
        [_paragraphs(c) for c in usable_cards],
        [[l for l in n.lines if l.strip()] for n in usable_news],
    )
    _card_tables(doc, usable_cards, bullet_numpr, size)
    _news_table(doc, usable_news, size)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
